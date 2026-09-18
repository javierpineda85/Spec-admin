"""Extrae una muestra histórica trazable. No escribe en Excel ni en la BD."""
from pathlib import Path
from collections import Counter, defaultdict
from difflib import SequenceMatcher
import hashlib
import json
import re
import unicodedata
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path('C:/Users/Usuario/Desktop/mayo 2025.xlsx')
OUT = ROOT / 'output/cronogramas_mayo_2025'
OUT.mkdir(parents=True, exist_ok=True)
db = json.loads(Path(__file__).with_name('catalogos.json').read_text(encoding='utf-8'))
wb = openpyxl.load_workbook(SOURCE, read_only=False, data_only=True)
ws = wb['MAYO 2025 (2)']
formula_wb = openpyxl.load_workbook(SOURCE, read_only=False, data_only=False)
formula_ws = formula_wb[ws.title]

def clean(value):
    return ' '.join(str(value or '').split())

def key(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', clean(value).upper()) if not unicodedata.combining(c))

def write_json(name, obj):
    (OUT / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

# Delimitaciones verificadas en el original; no depender de numeraciones manuales con #REF!.
blocks = [
    (2, 4, 6, 50, [18], 'equivalencia_por_nombre'),
    (58, 60, 62, 97, [29], 'equivalencia_por_nombre'),
    (107, 109, 111, 132, [27], 'equivalencia_por_nombre'),
    (141, 143, 145, 155, [19], 'equivalencia_por_nombre'),
    (162, 164, 166, 176, [2, 3, 5, 8, 9, 10, 11, 30], 'cuadro_agrupado_pendiente'),
    (191, 193, 195, 197, [16, 17], 'cuadro_agrupado_pendiente'),
    (208, 210, 212, 214, [6], 'equivalencia_por_nombre'),
    (219, 221, 223, 232, [20], 'equivalencia_por_nombre'),
    (236, 238, 240, 250, [], 'sin_equivalencia_segura'),
    (256, 258, 260, 265, [4], 'equivalencia_por_nombre'),
    (274, 276, 278, 279, [15], 'equivalencia_por_nombre'),
    (285, 287, 289, 294, [], 'sin_equivalencia_segura'),
    (300, 302, 304, 310, [14], 'equivalencia_por_nombre'),
]
dbobj = {int(o['idObjetivo']): o for o in db['objetivos']}
dbnames = defaultdict(list)
for u in db['usuarios']:
    dbnames[key(u['apellido'] + ' ' + u['nombre'])].append(u)

people = []
by_name = {}
objectives = []
roster = []
assignments = []
blanks = []
provenance = []
freq = Counter()
examples = defaultdict(list)
formula_inputs = []
for bi, (header, date_row, first, last, mapped, status) in enumerate(blocks, 1):
    oid = -bi
    days = [ws.cell(date_row, c).value for c in range(4, 35)]
    assert days == list(range(1, 32)), (header, days)
    title = clean(ws.cell(header, 3).value)
    assert title and clean(ws.cell(header + 1, 3).value).startswith('MAYO')
    objectives.append({'idObjetivo': oid, 'nombre': title, 'celda': f'C{header}',
                       'filas_nomina': [first, last], 'estado_correspondencia': status,
                       'objetivos_bd': [{'idObjetivo': i, 'nombre': dbobj[i]['nombre']} for i in mapped]})
    for row in ws.iter_rows(min_row=first, max_row=last, min_col=3, max_col=34):
        namecell, *cells = row
        rawname = namecell.value
        name = clean(rawname)
        assert name and not name.startswith('=')
        norm = key(name)
        if norm not in by_name:
            uid = -(len(people) + 1)
            by_name[norm] = uid
            matches = dbnames.get(norm, [])
            candidates = []
            if len(matches) != 1:
                # Un primer nombre abreviado no es una coincidencia exacta ni autoriza fusionar.
                source_tokens = Counter(norm.split())
                for u in db['usuarios']:
                    full_name = key(u['apellido'] + ' ' + u['nombre'])
                    if not (source_tokens - Counter(full_name.split())):
                        candidates.append({'idUsuario': u['idUsuario'], 'nombre': clean(u['apellido'] + ' ' + u['nombre']), 'criterio': 'nombre_abreviado_compatible_no_confirmado'})
                ranked = sorted(((SequenceMatcher(None, norm, n).ratio(), users) for n, users in dbnames.items()), key=lambda x: x[0], reverse=True)
                for score, users in ranked[:3]:
                    if score >= .78:
                        for u in users:
                            if not any(c['idUsuario'] == u['idUsuario'] for c in candidates):
                                candidates.append({'idUsuario': u['idUsuario'], 'nombre': clean(u['apellido'] + ' ' + u['nombre']), 'criterio': 'similitud_textual_no_confirmada', 'similitud': round(score, 3)})
            people.append({'idUsuario': uid, 'nombre_fuente': name, 'clave_comparacion': norm,
                           'variantes_fuente': [], 'celdas': [],
                           'estado_correspondencia': 'coincidencia_textual_unica' if len(matches) == 1 else ('coincidencia_textual_ambigua' if matches else 'sin_coincidencia_textual'),
                           'usuario_bd': ({k: matches[0][k] for k in ['idUsuario', 'nombre', 'apellido', 'activo', 'rol', 'categoria']} if len(matches) == 1 else None),
                           'candidatos_no_confirmados': candidates})
        uid = by_name[norm]
        person = people[-uid - 1]
        if rawname not in person['variantes_fuente']:
            person['variantes_fuente'].append(rawname)
        person['celdas'].append(namecell.coordinate)
        roster.append({'objetivo_id': oid, 'usuario_id': uid, 'celda_nombre': namecell.coordinate})
        for day, cell in enumerate(cells, 1):
            date = f'2025-05-{day:02d}'
            rawcode = cell.value
            code = clean(rawcode).upper()
            assert cell.data_type != 'e', (cell.coordinate, rawcode)
            if formula_ws[cell.coordinate].data_type == 'f':
                formula_inputs.append(cell.coordinate)
            if not code:
                blanks.append({'objetivo_id': oid, 'usuario_id': uid, 'fecha': date, 'celda': cell.coordinate})
                continue
            assert not code.startswith('=')
            assignments.append({'usuario_id': uid, 'objetivo_id': oid, 'fecha': date, 'codigo_turno': code})
            provenance.append({'indice_asignacion': len(assignments) - 1, 'celda': cell.coordinate, 'valor_original': rawcode})
            freq[code] += 1
            if len(examples[code]) < 4:
                examples[code].append(cell.coordinate)

# Destinos sugeridos por encabezados y códigos; no equivalencias históricas confirmadas.
refs = {'NOTT': [-1], 'PERR': [-2], 'BERM': [-9], 'METR': [-8], 'TUP': [-13],
        'H/8': [-10], 'C5,6': [-12], 'C/5,6': [-12], 'SALA': [-6], 'ETI': [-5], 'OFI': []}
services = {'D/GC': 5, 'D/GU': 8, 'D/LEM': 9, 'D/LH': 11, 'D/LUJ': 2, 'D/MA': 3, 'N/AR': 30, 'BE': 10}
absences = {'F': 'Franco', 'VAC': 'Vacaciones', 'P/EN': 'Parte de enfermo', 'F/INJ': 'Falta injustificada'}
legacy = {'6H','7H','8H','9H','9RF','9HEX','13H','14H','N15','D/LEM','D/GU','D/AR','D/LUJ','D/LH','D/GC','D/MA','BE'}
dictionary = []
for code, count in sorted(freq.items()):
    hours = None
    dest = []
    destdb = []
    status = 'no_aplica'
    if code in refs:
        kind = 'referencia_otro_objetivo'
        dest = refs[code]
        status = 'destino_sugerido_pendiente' if dest else 'destino_no_identificado'
        note = 'No contar como otro turno presencial. Verificar la asignación en el cuadro de destino.'
        if code in {'C5,6', 'C/5,6'}:
            note += ' Se conservan ambas grafías; la equivalencia específica aún no fue confirmada.'
    elif code in {'GP/D', 'GP/N'}:
        kind, hours, note = 'guardia_pasiva', 12, 'Disponibilidad sin presencia; conservar 12 h en el cálculo actual por instrucción del usuario.'
    elif code in absences:
        kind, hours, note = 'franco_o_ausencia', 0, absences[code]
    elif code in {'D', 'N'}:
        kind, hours, note = 'turno', 12, 'Entrada y salida dependen del puesto. No deducirlas de la duración.'
    elif re.fullmatch(r'\d+H', code):
        kind, hours, note = 'turno', int(code[:-1]), 'Duración indicada por el código; falta franja horaria.'
    elif code in {'N15','9RF','9HEX'}:
        kind, hours, note = 'turno_especial', int(re.search(r'\d+', code)[0]), 'Duración numérica; conservar significado especial y no inventar la franja.'
    elif code in services:
        kind, note = 'servicio_con_destino', 'Destino sugerido por código y catálogo actual. El cuadro DINAF-ETI agrupa destinos; falta confirmar duración y puesto históricos.'
        destdb = [services[code]]
        status = 'destino_sugerido_pendiente'
    else:
        raise AssertionError(f'Código sin clasificar: {code}')
    system_hours = 12 if code in {'D','N','GP/D','GP/N'} else int(code[:-1]) if re.fullmatch(r'\d+H', code) else 0
    visible = code in {'D','N','F','GP/D','GP/N','E','P','L','S'} or (re.fullmatch(r'\d+H',code) and 1 <= int(code[:-1]) <= 24)
    dictionary.append({'codigo': code, 'cantidad': count, 'categoria': kind,
                       'horas_muestra': hours, 'horas_codigo_actual_sin_sigla_configurada': system_hours,
                       'seleccionable_actualmente': bool(visible),
                       'reconocido_validador_global_heredado': code in legacy or code in {'F','GP/D','GP/N','E','P','S.','SALA','MIC','F/JUS','NOTT','GUE','PER','PAL','BOS','OFI'},
                       'objetivos_muestra_sugeridos': dest, 'objetivos_bd_sugeridos': destdb,
                       'estado_destino': status, 'nota': note, 'ejemplos': examples[code]})

by_person = defaultdict(set)
for r in roster:
    by_person[r['usuario_id']].add(r['objetivo_id'])
suspected_aliases = []
for i, p in enumerate(people):
    for q in people[i+1:]:
        a, b = p['clave_comparacion'], q['clave_comparacion']
        score = SequenceMatcher(None, a, b).ratio()
        if a.split()[0] == b.split()[0] and score >= .86:
            suspected_aliases.append({'usuario_a': p['idUsuario'], 'nombre_a': p['nombre_fuente'], 'usuario_b': q['idUsuario'], 'nombre_b': q['nombre_fuente'], 'estado': 'no_unificados_pendiente'})

dbv = {(int(x['usuario_id']),int(x['objetivo_id'])) for x in db['vigiladores']}
dbr = {(int(x['usuario_id']),int(x['objetivo_id'])) for x in db['referentes']}
links = []
for r in roster:
    p, o = people[-r['usuario_id']-1], objectives[-r['objetivo_id']-1]
    u = p['usuario_bd']
    if u and len(o['objetivos_bd']) == 1:
        pair = (int(u['idUsuario']), int(o['objetivos_bd'][0]['idObjetivo']))
        expected = dbv if u['categoria'] == 'operativo' else dbr if u['categoria'] == 'referente' else set()
        links.append({**r, 'idUsuario_bd': pair[0], 'objetivo_id_bd': pair[1], 'usuario_activo': u['activo'], 'categoria_actual': u['categoria'], 'vinculo_actual_segun_categoria': pair in expected})

summary = {
    'cuadros': len(objectives), 'filas_nomina': len(roster), 'identidades_textuales': len(people),
    'personas_en_varios_cuadros': sum(len(x)>1 for x in by_person.values()),
    'celdas_diarias': len(roster)*31, 'asignaciones_no_vacias': len(assignments), 'celdas_sin_asignacion': len(blanks),
    'codigos_distintos': len(dictionary), 'coincidencias_personal_bd': dict(Counter(p['estado_correspondencia'] for p in people)),
    'personas_sin_coincidencia_exacta_con_candidatos': sum(bool(p['candidatos_no_confirmados']) for p in people),
    'correspondencias_objetivos': dict(Counter(o['estado_correspondencia'] for o in objectives)),
    'categorias_celdas': dict(Counter(d['categoria'] for a in assignments for d in dictionary if d['codigo'] == a['codigo_turno'])),
    'codigos_no_seleccionables_actualmente': [d['codigo'] for d in dictionary if not d['seleccionable_actualmente']],
    'pares_de_nombres_por_confirmar': len(suspected_aliases), 'formulas_en_celdas_diarias': formula_inputs,
    'puestos_actuales': len(db['puestos']), 'horarios_actuales': len(db['horarios']), 'rotaciones_mayo_actuales': len(db['rotaciones_mayo'])
}
assert len(roster) == 173 and len(objectives) == 13
assert len(assignments) + len(blanks) == len(roster)*31
assert len(provenance) == len(assignments)
assert len({(a['usuario_id'],a['objetivo_id'],a['fecha']) for a in assignments}) == len(assignments)
assert sum(freq.values()) == len(assignments)
assert not formula_inputs, 'Revisar fórmulas antes de usar sus valores cacheados'
assert not {o['idObjetivo'] for o in objectives} & set(dbobj)
assert all(a['usuario_id'] in by_name.values() and a['objetivo_id'] in {o['idObjetivo'] for o in objectives} for a in assignments)
for a in assignments:
    assert set(a) == {'usuario_id','objetivo_id','fecha','codigo_turno'}

metadata = {'archivo': str(SOURCE), 'sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(), 'hoja': ws.title,
            'periodo': '2025-05', 'columnas_dias': 'D:AH', 'columna_nombre': 'C',
            'fecha_catalogos': db['capturado_en'], 'nota_catalogos': db['alcance'],
            'anotaciones_fuera_nomina': [{'celda': 'C269', 'texto': str(ws['C269'].value), 'tratamiento': 'Nota sin fechas ni turnos. Conservar como pendiente, no generar asignaciones.'}],
            'columnas_excluidas': {'A:B': 'Numeración y espacios; hay #REF! en B y empleados sin número.',
                                   'AI:AQ': 'Totales y liquidación histórica: no se incorporan al sistema ni se toman como resultado esperado.',
                                   'AR:CR': 'Fuera de la matriz diaria.'},
            'reglas': ['Identificadores negativos exclusivos de la muestra; nunca son IDs reales de la base.',
                       'No se corrigen grafías ni se unifican nombres parecidos sin evidencia suficiente.',
                       'Se igualan mayúsculas, tildes y espacios solo para coincidencia textual; no prueba identidad legal.',
                       'GP/D y GP/N conservan 12 horas; la disponibilidad no demuestra presencia.',
                       'No se inventan rol, tipo_turno, puesto ni horario que el Excel no identifica.',
                       'Se conservan códigos originales, incluyendo referencias. No importar todas las celdas como turnos trabajados.',
                       'No es un archivo de importación directa. Primero resolver correspondencias necesarias para cada prueba.']}
write_json('muestra_mayo_2025.json', {'fuente': metadata, 'resumen': summary,
                                    'objetivos': objectives, 'empleados': people, 'nomina_por_objetivo': roster,
                                    'asignaciones': assignments, 'sin_asignacion': blanks,
                                    'trazabilidad_asignaciones': provenance})
write_json('correspondencias_mayo_2025.json', {'fuente': metadata, 'codigos': dictionary,
                                           'nombres_pendientes': suspected_aliases, 'vinculos_contrastados': links,
                                           'horarios_actuales': {'puestos': db['puestos'], 'turnos': db['horarios'], 'rotaciones_mayo': db['rotaciones_mayo']}})

lines = ['# Muestra de cronogramas — mayo de 2025', '',
         'Primera parte: preparación y correspondencias. No se modificó el Excel, el código funcional ni la base de datos.', '',
         '## Alcance y reglas', '',
         '- Se preservan los 13 cuadros, sus nóminas y los 31 días. Las referencias permanecen diferenciadas en el diccionario de códigos.',
         '- GP/D y GP/N conservan **12 horas** por decisión del usuario. No se confunde esa suma con presencia efectiva.',
         '- Las columnas de liquidación, numeración y totales del Excel no se trasladan a la aplicación.',
         '- Los identificadores negativos pertenecen solo a esta muestra. No son IDs de la base ni autorizan una importación directa.',
         '- No se asignan roles o puestos históricos por suposición; los horarios actuales no prueban los de mayo de 2025.', '',
         '## Inventario', '',
         f"- {summary['filas_nomina']} filas de nómina; {summary['identidades_textuales']} nombres distintos tras normalizar espacios, tildes y mayúsculas.",
         f"- {summary['personas_en_varios_cuadros']} identidades textuales aparecen en varios cuadros; esto por sí solo no demuestra superposición.",
         f"- {summary['celdas_diarias']} celdas diarias: {len(assignments)} con código y {len(blanks)} sin asignación.",
         f"- {len(dictionary)} códigos distintos; {len(summary['codigos_no_seleccionables_actualmente'])} no están disponibles en el selector actual con el catálogo local leído.", '',
         '## Objetivos', '', '| Encabezado original | Personas | Correspondencia local | Estado |', '|---|---:|---|---|']
for o in objectives:
    count = o['filas_nomina'][1]-o['filas_nomina'][0]+1
    mapping = '; '.join(f"{x['nombre']} (ID {x['idObjetivo']})" for x in o['objetivos_bd']) or 'Sin coincidencia segura'
    lines.append(f"| {o['nombre']} ({o['celda']}) | {count} | {mapping} | {o['estado_correspondencia']} |")
lines += ['', 'Las equivalencias por nombre son correspondencias de catálogo, no una confirmación de asignaciones históricas. DINAF-ETI necesita distinguir los destinos por servicio; Palmira/Güemes necesita separar los dos destinos. No se crearon objetivos faltantes.', '',
          '## Empleados', '']
for state,count in summary['coincidencias_personal_bd'].items():
    lines.append(f'- {state}: {count}.')
lines += ['', 'Los candidatos similares quedan registrados para revisión y no se fusionan. Tampoco se da de alta a personal ausente del catálogo.', '',
          '| Nombre A | Nombre B | Tratamiento |', '|---|---|---|']
for p in suspected_aliases:
    lines.append(f"| {p['nombre_a']} | {p['nombre_b']} | Identificadores separados; pendiente confirmar |")
lines += ['', f"De los nombres sin coincidencia exacta, {summary['personas_sin_coincidencia_exacta_con_candidatos']} tienen candidatos por nombres abreviados o similitud. No debe interpretarse una falta de coincidencia exacta como ausencia del empleado en la base.",
          'La nota `GONZALO RODRIGO 8H` (C269) queda preservada por separado: no tiene fechas ni celdas de turno y no se convirtió en una asignación.']
lines += ['', f"Se pudieron contrastar {len(links)} vínculos persona–objetivo por coincidencia textual y objetivo individual. {sum(x['vinculo_actual_segun_categoria'] for x in links)} tienen vínculo según la categoría actual; {sum(not x['vinculo_actual_segun_categoria'] for x in links)} no. Estos datos actuales no invalidan el historial del Excel.", '',
          '## Diccionario de códigos', '', '| Código | Celdas | Clasificación | Horas muestra | Selector actual | Ejemplo |', '|---|---:|---|---:|---|---|']
for d in dictionary:
    lines.append(f"| {d['codigo']} | {d['cantidad']} | {d['categoria']} | {d['horas_muestra'] if d['horas_muestra'] is not None else 'Por confirmar / no aplica'} | {'Sí' if d['seleccionable_actualmente'] else 'No'} | {d['ejemplos'][0]} |")
lines += ['', 'La duración se conserva para D/N, guardias pasivas y códigos numéricos. VAC, P/EN y F/INJ conservan su texto histórico. No se transforman automáticamente a códigos modernos ni se crean siglas en la base.',
          'Las referencias NOTT/PERR/BERM/METR/TUP/H/8/C5,6/C/5,6/SALA/ETI/OFI no se cuentan como una segunda presencia. Sus destinos sugeridos y los códigos de servicio están detallados en correspondencias_mayo_2025.json.', '',
          '## Horarios disponibles y límites', '',
          f"El catálogo local contiene {len(db['puestos'])} puestos y {len(db['horarios'])} horarios, y no contiene rotaciones de puestos para mayo de 2025. Su detalle se guardó en las correspondencias.",
          'Los puestos actuales pertenecen a Parque Central y Empresa Marsan, fuera de los cuadros de la muestra; no aportan franjas horarias aplicables a estos objetivos.',
          'El Excel no identifica entrada/salida ni puesto por celda. No se puede confirmar ausencia o presencia de superposición horaria con la duración únicamente.', '',
          '## Controles de preparación', '',
          '- Se verificaron las 13 cabeceras de fechas: días 1 a 31 en D:AH.',
          '- Se incluyeron empleados sin numeración y filas cuya numeración contiene #REF!. La matriz diaria no contiene errores de Excel ni fórmulas.',
          '- Cada asignación tiene celda de origen y valor original; cada blanco se conserva como ausencia de asignación.',
          '- Se verificó la unicidad usuario–cuadro–fecha, las referencias internas y la suma celdas con código + celdas vacías.',
          '- Los datos de asignación usan solo usuario_id, objetivo_id, fecha y codigo_turno. Los datos de trazabilidad pertenecen al material de prueba; no son columnas nuevas del sistema.', '',
          '## Archivos', '',
          '- `muestra_mayo_2025.json`: fuente, nóminas, asignaciones y trazabilidad.',
          '- `correspondencias_mayo_2025.json`: códigos, candidatos pendientes, vínculos y horarios actuales.',
          '', '## Pendientes para las siguientes partes', '',
          '- Confirmar identidades de los nombres parecidos antes de unirlos en pruebas entre objetivos.',
          '- Resolver los cuadros agrupados y las referencias sin destino inequívoco; conservar C5,6 y C/5,6 como variantes pendientes.',
          '- Elegir puestos/horarios explícitos en los escenarios de prueba. No usar horarios inventados como evidencia histórica.',
          '- Validar compatibilidad de los códigos históricos antes de cargar la muestra mediante la interfaz.',
          '', f"Fuente: {SOURCE} — hoja `{ws.title}`. SHA-256: `{metadata['sha256']}`.",
          f"Catálogos locales consultados: {db['capturado_en']}. Sin escrituras en base de datos."]
(OUT/'revision_muestra.md').write_text('\n'.join(lines)+'\n', encoding='utf-8')
print(json.dumps(summary, ensure_ascii=False, indent=2))
print('PENDING_NAMES', json.dumps(suspected_aliases, ensure_ascii=False))
